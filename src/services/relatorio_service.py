"""Serviço de geração de relatórios Word."""

from __future__ import annotations

import io
from docx import Document
from docx.shared import Inches


class RelatorioService:
    """Gera relatórios DOCX executivos simples."""

    def gerar_relatorio_inspecao(self, status: str, caminhos: list[str]) -> bytes:
        """Inclui todas as evidências persistidas no laudo da inspeção."""
        doc = Document()
        doc.add_heading("Laudo de inspeção", level=1)
        doc.add_paragraph(f"Status geral: {status}")
        doc.add_heading("Evidências fotográficas", level=2)
        for index, caminho in enumerate(caminhos, 1):
            doc.add_paragraph(f"Evidência {index}")
            doc.add_picture(caminho, width=Inches(5))
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def gerar_relatorio_ocorrencia(self, titulo: str, secoes: dict[str, str]) -> bytes:
        """Gera conteúdo de relatório em bytes."""
        doc = Document()
        doc.add_heading(titulo, level=1)
        for secao, conteudo in secoes.items():
            doc.add_heading(secao, level=2)
            doc.add_paragraph(conteudo)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
