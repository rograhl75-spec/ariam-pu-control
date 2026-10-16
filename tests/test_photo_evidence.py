from io import BytesIO
import base64
from pathlib import Path
from shutil import copyfile
import json
import sqlite3
from types import SimpleNamespace
from unittest.mock import Mock
from zipfile import ZipFile

from docx import Document
from PIL import Image
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest
from streamlit.proto.Common_pb2 import FileURLs
from streamlit.runtime.uploaded_file_manager import UploadedFile, UploadedFileRec

from ui import photo_evidence, rear_camera
from ui.photo_evidence import CAMERA_SOURCE, FILE_SOURCE, prepare_photo, select_photos


ROOT = Path(__file__).resolve().parents[1]


def make_upload(name="camera.jpg", fmt="JPEG", color="red"):
    buffer = BytesIO()
    Image.new("RGB", (8, 4), color=color).save(buffer, format=fmt)
    content = buffer.getvalue()
    return UploadedFile(
        UploadedFileRec(file_id=name, name=name, type=Image.MIME[fmt], data=content),
        FileURLs(),
    )


def mock_capture(monkeypatch, capture):
    monkeypatch.setattr(rear_camera, "_component", lambda **kwargs: {
        "session": kwargs["session"],
        "status": "photo" if capture else "empty",
        "data": "data:image/jpeg;base64," + base64.b64encode(capture.getvalue()).decode() if capture else None,
    })


@pytest.mark.parametrize("fmt, name", [("JPEG", "camera.jpg"), ("PNG", "arquivo.png"), ("WEBP", "foto.webp")])
def test_prepare_photo_normalizes_capture_and_upload(fmt, name):
    upload = make_upload(name, fmt)
    photo = prepare_photo(upload)
    assert photo.name == f"{Path(name).stem}.png"
    with Image.open(BytesIO(photo.content)) as image:
        assert image.format == "PNG"
        assert image.size == (8, 4)
    doc = Document()
    doc.add_picture(BytesIO(photo.content))
    assert len(doc.inline_shapes) == 1


def test_prepare_photo_applies_camera_orientation():
    buffer = BytesIO()
    image = Image.new("RGB", (8, 4))
    exif = image.getexif()
    exif[274] = 6
    image.save(buffer, format="JPEG", exif=exif)
    photo = prepare_photo(SimpleNamespace(name="camera.jpg", getvalue=buffer.getvalue))
    with Image.open(BytesIO(photo.content)) as normalized:
        assert normalized.size == (4, 8)


def test_prepare_photo_preserves_transparent_png_evidence():
    image = Image.new("RGBA", (2, 1), (0, 0, 0, 0))
    image.putpixel((1, 0), (0, 0, 0, 255))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    photo = prepare_photo(SimpleNamespace(name="evidencia.png", getvalue=buffer.getvalue))
    with Image.open(BytesIO(photo.content)) as normalized:
        assert normalized.getpixel((0, 0)) == (0, 0, 0, 0)
        assert normalized.getpixel((1, 0)) == (0, 0, 0, 255)
    document = Document()
    document.add_picture(BytesIO(photo.content))
    assert len(document.inline_shapes) == 1


@pytest.mark.parametrize("content", [b"", b"not-an-image", b"x" * (5 * 1024 * 1024 + 1)])
def test_prepare_photo_rejects_invalid_or_oversized_bytes(content):
    with pytest.raises(ValueError):
        prepare_photo(SimpleNamespace(name="foto.png", getvalue=lambda: content))


def test_source_switch_never_falls_back_to_other_source():
    upload, capture = make_upload(), make_upload()
    assert select_photos(FILE_SOURCE, [upload], capture) == [upload]
    assert select_photos(CAMERA_SOURCE, [upload], capture) == [capture]
    assert select_photos(CAMERA_SOURCE, [upload], None) == []
    assert select_photos(FILE_SOURCE, [], capture) == []
    assert select_photos(FILE_SOURCE, [], None) == []
    with pytest.raises(ValueError):
        select_photos("invalid", [upload], capture)


def test_render_invalid_active_photo_blocks_save(monkeypatch):
    ui = Mock()
    ui.radio.return_value = CAMERA_SOURCE
    ui.file_uploader.return_value = make_upload()
    monkeypatch.setattr(photo_evidence, "render_rear_camera", lambda *a, **k: (
        SimpleNamespace(name="camera.jpg", getvalue=lambda: b"invalid")
    ))
    monkeypatch.setattr(photo_evidence, "st", ui)
    assert photo_evidence.render_photo_evidence("test") is None
    ui.error.assert_called_once()


def test_camera_component_and_uploader_render():
    app = AppTest.from_string(
        "from ui.photo_evidence import render_photo_evidence\n"
        "render_photo_evidence('photo')\n"
    ).run()
    assert not app.exception
    assert len(app.get("component_instance")) == 1
    assert len(app.get("file_uploader")) == 1
    assert CAMERA_SOURCE in app.radio[0].options
    assert not app.get("file_uploader")[0].proto.disabled
    app.radio[0].set_value(CAMERA_SOURCE).run()
    assert not app.exception
    assert app.get("file_uploader")[0].proto.disabled


@pytest.mark.parametrize("source", [FILE_SOURCE, CAMERA_SOURCE])
def test_modular_saves_only_active_source_names(tmp_path, monkeypatch, source):
    import streamlit as st
    from src.pages import inspecoes

    db_path = tmp_path / "inspections.db"
    monkeypatch.setattr(inspecoes, "get_settings", lambda: SimpleNamespace(db_path=db_path))
    uploads = [make_upload("primeira.jpg"), make_upload("segunda.png", "PNG")]
    capture = make_upload("camera.jpg")
    monkeypatch.setattr(st, "file_uploader", lambda *a, **k: uploads)
    mock_capture(monkeypatch, capture)
    app = AppTest.from_string("from src.pages.inspecoes import render\nrender()").run()
    app.radio(key="inspecao_source").set_value(source).run()
    next(button for button in app.button if button.label == "Salvar inspeção").click().run()
    assert not app.exception
    with sqlite3.connect(db_path) as connection:
        payload = json.loads(connection.execute("SELECT data_json FROM measurements").fetchone()[0])
    selected = uploads if source == FILE_SOURCE else [capture]
    assert payload["fotos"] == [prepare_photo(photo).name for photo in selected]


@pytest.mark.parametrize("source", [FILE_SOURCE, CAMERA_SOURCE])
def test_legacy_form_preview_save_and_word(tmp_path, monkeypatch, source):
    import docx
    import requests
    import streamlit as st

    monkeypatch.chdir(tmp_path)
    base = "Relatorio_Processo_Injecao_Atualizado.xlsx"
    copyfile(ROOT / base, tmp_path / base)
    monkeypatch.setattr(requests, "get", lambda *a, **k: SimpleNamespace(
        json=lambda: {"current_weather": {"temperature": 24}}
    ))
    upload = make_upload("arquivo.png", "PNG", "blue")
    capture = make_upload()
    monkeypatch.setattr(st, "file_uploader", lambda *a, **k: upload)
    mock_capture(monkeypatch, capture)
    documents = []

    def document_factory():
        document = Document()
        documents.append(document)
        return document

    monkeypatch.setattr(docx, "Document", document_factory)
    app = AppTest.from_file(str(ROOT / "app_tablet.py"), default_timeout=20)
    app.session_state["senha_correta"] = True
    app.session_state["menu_ativo"] = "inspecoes"
    app.run()
    assert not app.exception
    next(field for field in app.text_area if field.label == "Observações de Manutenção:").set_value("Mantido na prévia")
    app.run()
    app.radio(key="foto_insp_source").set_value(source).run()
    assert next(field for field in app.text_area if field.label == "Observações de Manutenção:").value == "Mantido na prévia"
    next(button for button in app.button if button.label == "Conferir prévia da foto").click().run()
    assert not app.exception
    assert not (tmp_path / "Log_Inspecoes_Semanais.xlsx").exists()
    assert next(field for field in app.text_area if field.label == "Observações de Manutenção:").value == "Mantido na prévia"
    next(button for button in app.button if button.label == "Salvar Inspeção e Gerar Laudo Técnico").click().run()
    assert not app.exception
    row = pd.read_excel(tmp_path / "Log_Inspecoes_Semanais.xlsx").iloc[0]
    selected = capture if source == CAMERA_SOURCE else upload
    assert (tmp_path / row["Foto_Path"]).read_bytes() == prepare_photo(selected).content
    assert row["Manutenção Obs"] == "Mantido na prévia"
    report = BytesIO()
    documents[-1].save(report)
    with ZipFile(report) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert len(media) == 1
        assert archive.read(media[0]) == prepare_photo(selected).content

    # A foto antiga da outra origem não deve reaparecer ao salvar sem imagem.
    if source == CAMERA_SOURCE:
        monkeypatch.setattr(st, "file_uploader", lambda *a, **k: None)
        app.radio(key="foto_insp_source").set_value(FILE_SOURCE)
    else:
        mock_capture(monkeypatch, None)
        app.radio(key="foto_insp_source").set_value(CAMERA_SOURCE)
    next(button for button in app.button if button.label == "Salvar Inspeção e Gerar Laudo Técnico").click().run()
    assert not app.exception
    rows = pd.read_excel(tmp_path / "Log_Inspecoes_Semanais.xlsx")
    assert len(rows) == 2
    assert pd.isna(rows.iloc[1]["Foto_Path"])
